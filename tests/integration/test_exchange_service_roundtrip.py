"""Round-trip: импорт собственного экспорта не должен ломать machine values."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

from app.application.services.exchange_service import ExchangeService
from scripts.seed_demo_data import seed
from tests.integration.test_exchange_service_import_reports import (
    make_session_factory,
    seed_actor,
)


def _seed_roundtrip_db(tmp_path: Path, db_name: str) -> tuple[ExchangeService, int]:
    session_factory = make_session_factory(tmp_path / db_name)
    actor_id = seed_actor(session_factory)
    with session_factory() as session:
        seed(session, clear=True, id_store_path=tmp_path / f"{db_name}.ids.json")
    return ExchangeService(session_factory=session_factory), actor_id


def _read_sheet_rows(workbook_path: Path, sheet_name: str) -> list[dict[str, object]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    return [
        dict(zip(headers, row, strict=False))
        for row in sheet.iter_rows(min_row=2, values_only=True)
    ]


def test_excel_roundtrip_no_errors(tmp_path: Path) -> None:
    """Импорт собственного Excel-экспорта не должен давать CHECK-ошибок."""
    svc, actor_id = _seed_roundtrip_db(tmp_path, "rt.db")
    export_path = tmp_path / "export.xlsx"

    svc.export_excel(export_path, exported_by="rt_admin", actor_id=actor_id)
    result = svc.import_excel(export_path, actor_id=actor_id, mode="merge")

    assert result["error_count"] == 0, (
        f"{result['error_count']} ошибок при round-trip:\n"
        + "\n".join(str(error) for error in result["errors"][:10])
    )


def test_csv_roundtrip_no_errors(tmp_path: Path) -> None:
    """Импорт собственного CSV не должен давать ошибок."""
    svc, actor_id = _seed_roundtrip_db(tmp_path, "rt_csv.db")

    for table in ("patients", "lab_sample", "sanitary_sample", "emr_case"):
        csv_path = tmp_path / f"{table}.csv"
        svc.export_csv(csv_path, table, actor_id=actor_id)
        result = svc.import_csv(csv_path, table, actor_id=actor_id, mode="merge")
        assert result["error_count"] == 0, (
            f"CSV {table}: {result['error_count']} ошибок: "
            + str(result.get("errors", [])[:5])
        )


def test_excel_export_preserves_machine_enum_values(tmp_path: Path) -> None:
    """Excel должен хранить machine values, пригодные для обратного импорта."""
    svc, actor_id = _seed_roundtrip_db(tmp_path, "rt_enum.xlsx.db")
    export_path = tmp_path / "enum.xlsx"

    svc.export_excel(export_path, exported_by="rt_admin", actor_id=actor_id)
    workbook = load_workbook(export_path, read_only=True, data_only=True)
    lab_sheet = workbook["Лабораторные пробы"]
    lab_headers = [cell.value for cell in next(lab_sheet.iter_rows(min_row=1, max_row=1))]
    lab_values = [
        dict(zip(lab_headers, row, strict=False))
        for row in lab_sheet.iter_rows(min_row=2, values_only=True)
    ]

    assert "primary" in {row.get("Тип исследования") for row in lab_values}
    assert "первичное" not in {row.get("Тип исследования") for row in lab_values}
    assert "valid" in {row.get("Статус QC") for row in lab_values}
    assert "валидный" not in {row.get("Статус QC") for row in lab_values}
    assert {row.get("Результат роста") for row in lab_values} <= {0, 1, None}


def test_excel_export_localized_true_uses_readable_enum_values(tmp_path: Path) -> None:
    svc, actor_id = _seed_roundtrip_db(tmp_path, "rt_enum_localized.xlsx.db")
    export_path = tmp_path / "enum_localized.xlsx"

    svc.export_excel(export_path, exported_by="rt_admin", actor_id=actor_id, localized=True)

    patient_rows = _read_sheet_rows(export_path, "Пациенты")
    lab_rows = _read_sheet_rows(export_path, "Лабораторные пробы")
    version_rows = _read_sheet_rows(export_path, "Версии ЭМЗ")
    diagnosis_rows = _read_sheet_rows(export_path, "Диагнозы ЭМЗ")
    susceptibility_rows = _read_sheet_rows(export_path, "Лаб. чувствительность")
    form100_rows = _read_sheet_rows(export_path, "Форма 100")

    assert "М" in {row.get("Пол") for row in patient_rows}
    assert "M" not in {row.get("Пол") for row in patient_rows}
    assert "Первичное" in {row.get("Тип исследования") for row in lab_rows}
    assert "primary" not in {row.get("Тип исследования") for row in lab_rows}
    assert "Действителен" in {row.get("Статус QC") for row in lab_rows}
    assert "valid" not in {row.get("Статус QC") for row in lab_rows}
    assert "Рост выявлен" in {row.get("Результат роста") for row in lab_rows}
    assert "Выписан" in {row.get("Тип исхода") for row in version_rows}
    assert "Средняя" in {row.get("Тяжесть") for row in version_rows}
    assert "При поступлении" in {row.get("Тип диагноза") for row in diagnosis_rows}
    assert "Чувствительный" in {row.get("RIS") for row in susceptibility_rows}
    assert "Диско-диффузионный" in {row.get("Метод") for row in susceptibility_rows}
    assert "Черновик" in {row.get("Статус") for row in form100_rows}
    assert "DRAFT" not in {row.get("Статус") for row in form100_rows}


def test_excel_localized_roundtrip_restores_machine_enum_values(tmp_path: Path) -> None:
    session_factory = make_session_factory(tmp_path / "rt_enum_localized_import.db")
    actor_id = seed_actor(session_factory)
    with session_factory() as session:
        seed(session, clear=True, id_store_path=tmp_path / "rt_enum_localized_import.ids.json")
    svc = ExchangeService(session_factory=session_factory)
    export_path = tmp_path / "enum_localized_roundtrip.xlsx"

    svc.export_excel(export_path, exported_by="rt_admin", actor_id=actor_id, localized=True)
    result = svc.import_excel(export_path, actor_id=actor_id, mode="merge")

    assert result["error_count"] == 0, (
        f"{result['error_count']} ошибок при localized round-trip:\n"
        + "\n".join(str(error) for error in result["errors"][:10])
    )
    with session_factory() as session:
        bad_sex = session.execute(
            text("SELECT COUNT(*) FROM patients WHERE sex IS NOT NULL AND sex NOT IN ('M','F','U')")
        ).scalar_one()
        bad_study = session.execute(
            text(
                "SELECT COUNT(*) FROM lab_sample "
                "WHERE study_kind IS NOT NULL "
                "AND study_kind NOT IN ('primary','repeat')"
            )
        ).scalar_one()
        bad_outcome = session.execute(
            text(
                "SELECT COUNT(*) FROM emr_case_version "
                "WHERE outcome_type IS NOT NULL "
                "AND outcome_type NOT IN ('discharge','transfer','death')"
            )
        ).scalar_one()
        bad_kind = session.execute(
            text(
                "SELECT COUNT(*) FROM emr_diagnosis "
                "WHERE kind IS NOT NULL "
                "AND kind NOT IN ('admission','discharge','complication')"
            )
        ).scalar_one()
        bad_form100_status = session.execute(
            text("SELECT COUNT(*) FROM form100 WHERE status IS NOT NULL AND status NOT IN ('DRAFT','SIGNED')")
        ).scalar_one()

    assert bad_sex == 0
    assert bad_study == 0
    assert bad_outcome == 0
    assert bad_kind == 0
    assert bad_form100_status == 0


def test_csv_export_localized_true_roundtrips_readable_enum_values(tmp_path: Path) -> None:
    svc, actor_id = _seed_roundtrip_db(tmp_path, "rt_csv_localized.db")
    csv_path = tmp_path / "lab_sample_localized.csv"

    svc.export_csv(csv_path, "lab_sample", actor_id=actor_id, localized=True)

    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert "Первичное" in {row.get("Тип исследования") for row in rows}
    assert "primary" not in {row.get("Тип исследования") for row in rows}
    assert "Действителен" in {row.get("Статус QC") for row in rows}
    assert "Рост выявлен" in {row.get("Результат роста") for row in rows}

    result = svc.import_csv(csv_path, "lab_sample", actor_id=actor_id, mode="merge")

    assert result["error_count"] == 0, str(result.get("errors", [])[:5])


def test_excel_roundtrip_keeps_database_enum_values_valid(tmp_path: Path) -> None:
    """После export->import enum-поля должны остаться машинными."""
    session_factory = make_session_factory(tmp_path / "rt_enum.db")
    actor_id = seed_actor(session_factory)
    with session_factory() as session:
        seed(session, clear=True, id_store_path=tmp_path / "rt_enum.ids.json")
    svc = ExchangeService(session_factory=session_factory)
    export_path = tmp_path / "enum_roundtrip.xlsx"

    svc.export_excel(export_path, exported_by="rt_admin", actor_id=actor_id)
    svc.import_excel(export_path, actor_id=actor_id, mode="merge")

    with session_factory() as session:
        bad_study = session.execute(
            text(
                "SELECT COUNT(*) FROM lab_sample "
                "WHERE study_kind IS NOT NULL "
                "AND study_kind NOT IN ('primary','repeat')"
            )
        ).scalar_one()
        bad_kind = session.execute(
            text(
                "SELECT COUNT(*) FROM emr_diagnosis "
                "WHERE kind IS NOT NULL "
                "AND kind NOT IN ('admission','discharge','complication')"
            )
        ).scalar_one()
    assert bad_study == 0, f"{bad_study} строк lab_sample с переведённым study_kind"
    assert bad_kind == 0, f"{bad_kind} строк emr_diagnosis с переведённым kind"
